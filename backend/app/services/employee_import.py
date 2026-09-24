"""Bounded CSV/XLSX parsing and shared preview/commit validation. Never creates logins."""

import csv
import io
import re
import zipfile
from collections import Counter
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select
from ..models import Employee

HEADERS = ["First Name", "Last Name", "Email", "Employee ID", "Manager", "Manager Eligible"]
MAX_ROWS = 1000


def parse(content, filename):
    try:
        if filename.lower().endswith(".csv"):
            values = csv.reader(io.StringIO(content.decode("utf-8-sig")), strict=True)
        elif filename.lower().endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 200 or sum(v.file_size for v in archive.infolist()) > 20_000_000:
                    raise ValueError("expanded workbook too large")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            if len(workbook.worksheets) != 1:
                workbook.close()
                raise ValueError("one sheet required")
            values = workbook.active.iter_rows(values_only=True)
        else:
            raise HTTPException(422, "Use the CSV or XLSX employee template")
        result = []
        try:
            headers = [str(x or "").strip() for x in next(values)]
            if (
                len(headers) > 6
                or len(set(headers)) != len(headers)
                or not {"First Name", "Last Name"}.issubset(headers)
                or any(h not in HEADERS for h in headers)
            ):
                raise ValueError("headers")
            for index, row in enumerate(values):
                if index >= MAX_ROWS:
                    raise ValueError("rows")
                if len(row) > len(headers) and any(x not in (None, "") for x in row[len(headers) :]):
                    raise ValueError("columns")
                cells = [str(x if x is not None else "").strip() for x in row[: len(headers)]]
                if not any(cells):
                    continue
                if any(len(c) > 254 or c.startswith("=") for c in cells):
                    raise ValueError("cell")
                result.append({"row": index + 2, **dict(zip(headers, cells))})
        finally:
            if filename.lower().endswith(".xlsx"):
                workbook.close()
        if not result:
            raise ValueError("empty")
        return result
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            422,
            "Invalid import. Use one sheet, template headers, UTF-8 CSV or XLSX, no formulas, at most 1,000 rows and 254 characters per cell.",
        ) from None


def validate(rows, db, organization_id):
    existing = db.scalars(select(Employee).where(Employee.organization_id == organization_id)).all()
    identifiers = {e.external_id: e for e in existing if e.external_id}
    emails = {e.email: e for e in existing if e.email}
    names = Counter(e.name.casefold() for e in existing)
    records = []
    for raw in rows:
        first, last = raw.get("First Name", ""), raw.get("Last Name", "")
        name = (first + " " + last).strip()
        email = raw.get("Email", "").lower() or None
        external_id = raw.get("Employee ID", "").casefold() or None
        eligibility = raw.get("Manager Eligible", "").casefold()
        errors, warnings = [], []
        if not first or not last or len(name) > 120:
            errors.append("First and last name are required; combined name must be at most 120 characters")
        if email and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
            errors.append("Invalid email")
        if external_id and len(external_id) > 100:
            errors.append("Employee ID must be at most 100 characters")
        if eligibility not in {"", "true", "false", "yes", "no", "1", "0"}:
            errors.append("Manager Eligible must be yes or no")
        records.append(
            {
                "row": raw["row"],
                "name": name,
                "email": email,
                "external_id": external_id,
                "manager_eligible": eligibility in {"true", "yes", "1"},
                "manager": raw.get("Manager", "").casefold(),
                "errors": errors,
                "warnings": warnings,
                "manager_key": None,
            }
        )
        names[name.casefold()] += 1
    ids = Counter(r["external_id"] for r in records if r["external_id"])
    addresses = Counter(r["email"] for r in records if r["email"])
    graph = {}
    for i, r in enumerate(records):
        if r["external_id"] and (r["external_id"] in identifiers or ids[r["external_id"]] > 1):
            r["errors"].append("Duplicate employee ID (file or existing organization)")
        if r["email"] and (r["email"] in emails or addresses[r["email"]] > 1):
            r["errors"].append("Duplicate email (file or existing organization)")
        if names[r["name"].casefold()] > 1:
            r["warnings"].append("Another employee has this name. Confirm this is a different person")
        key = r["manager"]
        if key:
            matches = {("existing", e.id): e for e in existing if key in {e.external_id, e.email}}
            matches.update(
                {("new", j): other for j, other in enumerate(records) if key in {other["external_id"], other["email"]}}
            )
            if len(matches) != 1:
                r["errors"].append("Manager identifier/email was not found or is ambiguous in this organization/import")
            else:
                target, manager = next(iter(matches.items()))
                eligible = (
                    manager.manager_eligible and manager.active
                    if target[0] == "existing"
                    else manager["manager_eligible"]
                )
                if not eligible:
                    r["errors"].append("Manager must be active and explicitly Manager Eligible")
                elif target == ("new", i):
                    r["errors"].append("Employee cannot manage themselves")
                else:
                    r["manager_key"] = list(target)
                    if target[0] == "new":
                        graph[i] = target[1]
    for i, r in enumerate(records):
        seen, node = set(), i
        while node in graph:
            if node in seen:
                r["errors"].append("Reporting relationships contain a cycle")
                break
            seen.add(node)
            node = graph[node]
    return {
        "total": len(records),
        "ready": sum(not r["errors"] for r in records),
        "errors": sum(bool(r["errors"]) for r in records),
        "warnings": sum(bool(r["warnings"]) for r in records),
        "rows": records,
    }
